"""验 pick rank 端到端：归集 → 排除过滤 → 打分 → 四份产物落盘。

用合成快照，不触网。关键是验证「成交速率需要 ≥2 个不同时刻的快照」这条硬约束
以及被过滤清单的可复核性。
"""
from __future__ import annotations

import csv
import json
from pathlib import Path

from goofish_cli.commands.pick import rank as rank_cmd

T1 = "2026-01-01T10:00:00+08:00"
T2 = "2026-01-02T10:00:00+08:00"


def _item(ts: str, item_id: str, keyword: str, title: str, price: float, want: int,
          uid: str, nick: str, rate_count: int) -> dict:
    return {
        "type": "item", "ts": ts, "item_id": item_id, "keyword": keyword,
        "keywords": [keyword], "title": title, "price": f"¥{price:g}",
        "price_num": price, "want": want, "seller_uid": uid, "seller_nick": nick,
        "seller_rate_count": rate_count, "seller_good_rate": 90, "area": "北京",
        "publish_ts": 1, "cat_id": "5001", "c_cat_id": "6001", "tb_cat_id": "7001",
        "tag": "freeship", "item_type": "goods",
    }


def _seller(ts: str, uid: str, nick: str, rate_count: int, keyword: str) -> dict:
    return {
        "type": "seller", "ts": ts, "uid": uid, "nick": nick,
        "rate_count": rate_count, "good_rate": 90, "items_sampled": 1,
        "keywords": [keyword],
    }


def _write_snapshot(root: Path, name: str, records: list[dict]) -> None:
    (root / "snapshots").mkdir(parents=True, exist_ok=True)
    path = root / "snapshots" / name
    with path.open("w", encoding="utf-8") as fh:
        for rec in records:
            fh.write(json.dumps(rec, ensure_ascii=False) + "\n")


def _fixture(root: Path) -> None:
    # 快照 1
    _write_snapshot(root, "20260101-100000.jsonl", [
        {"type": "keyword", "ts": T1, "keyword": "考研 资料", "supply": 5000,
         "fetched": 3, "secs": 1.0, "error": ""},
        {"type": "keyword", "ts": T1, "keyword": "设计素材", "supply": 800_000,
         "fetched": 1, "secs": 1.0, "error": ""},
        _item(T1, "1", "考研 资料", "2026考研全套资料 电子版", 9.0, 100, "u1", "甲", 1000),
        _item(T1, "2", "考研 资料", "考研真题 笔记合集", 5.0, 50, "u2", "乙", 200),
        _item(T1, "3", "考研 资料", "腾讯视频会员日卡直充", 1.0, 999, "u3", "丙", 100),
        _item(T1, "4", "设计素材", "别墅庭院设计素材 源文件", 2.0, 10, "u4", "丁", 100),
        _seller(T1, "u1", "甲", 1000, "考研 资料"),
        _seller(T1, "u2", "乙", 200, "考研 资料"),
        _seller(T1, "u3", "丙", 100, "考研 资料"),
        _seller(T1, "u4", "丁", 100, "设计素材"),
    ])
    # 快照 2：卖家评价数增长 → 成交速率可算
    _write_snapshot(root, "20260102-100000.jsonl", [
        {"type": "keyword", "ts": T2, "keyword": "考研 资料", "supply": 5000,
         "fetched": 3, "secs": 1.0, "error": ""},
        {"type": "keyword", "ts": T2, "keyword": "设计素材", "supply": 800_000,
         "fetched": 1, "secs": 1.0, "error": ""},
        _item(T2, "1", "考研 资料", "2026考研全套资料 电子版", 9.0, 120, "u1", "甲", 1002),
        _item(T2, "2", "考研 资料", "考研真题 笔记合集", 5.0, 55, "u2", "乙", 205),
        _item(T2, "3", "考研 资料", "腾讯视频会员日卡直充", 1.0, 1000, "u3", "丙", 100),
        _item(T2, "4", "设计素材", "别墅庭院设计素材 源文件", 2.0, 12, "u4", "丁", 100),
        _seller(T2, "u1", "甲", 1002, "考研 资料"),
        _seller(T2, "u2", "乙", 205, "考研 资料"),
        _seller(T2, "u3", "丙", 100, "考研 资料"),
        _seller(T2, "u4", "丁", 100, "设计素材"),
    ])


def _weights() -> dict:
    return {
        "dimensions": {
            "imbalance": {"weight": 0.3},
            "delta_rate": {"weight": 0.25},
            "want": {"weight": 0.15},
            "seller_base": {"weight": 0.15},
            "samples": {"weight": 0.1},
            "price": {"weight": 0.05},
        },
        "exclusion": {
            "renewal_strong": ["日卡", "直充"],
            "renewal_weak": ["会员"],
            "human": ["代练"],
            "low_price_ceiling": 5.0,
        },
        "min_samples": 1,
    }


def test_build_aggregates_and_excludes(tmp_path):
    _fixture(tmp_path)
    result = rank_cmd._build(tmp_path, _weights())

    by_kw = {row["keyword"]: row for row in result["categories"]}
    assert set(by_kw) == {"考研 资料", "设计素材"}

    kaoyan = by_kw["考研 资料"]
    assert kaoyan["samples"] == 2                       # 被排除那条不计入
    assert kaoyan["want_median"] == 87.5                # median(120, 55)
    assert kaoyan["excluded_ratio"] == round(1 / 3, 3)  # 3 条里排除 1 条
    assert kaoyan["supply_capped"] is False
    assert kaoyan["imbalance_raw"] is not None

    # 卖家成交速率：u1 2/天、u2 5/天（u3 被排除不影响）→ 中位 3.5
    assert kaoyan["delta_rate_median"] == 3.5


def test_capped_supply_disables_imbalance(tmp_path):
    _fixture(tmp_path)
    result = rank_cmd._build(tmp_path, _weights())
    sheji = next(r for r in result["categories"] if r["keyword"] == "设计素材")
    assert sheji["supply_capped"] is True
    assert sheji["imbalance_raw"] is None, "供给被截断时失衡无法计算，不能给假值"
    assert "imbalance" in sheji["dims_degraded"]


def test_excluded_items_carry_reason(tmp_path):
    _fixture(tmp_path)
    result = rank_cmd._build(tmp_path, _weights())
    excluded = [r for r in result["items"] if r["excluded"]]
    assert len(excluded) == 1
    assert excluded[0]["item_id"] == "3"
    assert "续费型" in excluded[0]["exclude_reason"]


def test_seller_rows_ranked_by_delta(tmp_path):
    _fixture(tmp_path)
    result = rank_cmd._build(tmp_path, _weights())
    sellers = {row["uid"]: row for row in result["sellers"]}
    assert sellers["u2"]["delta_rate"] == 5.0
    assert sellers["u1"]["delta_rate"] == 2.0
    assert result["sellers"][0]["uid"] == "u2", "卖家榜应按成交速率降序"


def test_single_snapshot_degrades_delta_dimension(tmp_path):
    """只有一份快照时成交速率算不出来 —— 该维度必须降级而不是当 0 分。"""
    _fixture(tmp_path)
    (tmp_path / "snapshots" / "20260102-100000.jsonl").unlink()
    result = rank_cmd._build(tmp_path, _weights())
    for row in result["categories"]:
        assert row["delta_rate_raw"] is None
        assert "delta_rate" in row["dims_degraded"]
        assert "delta_rate" not in row["dims_used"]
        assert row["score"] is not None, "剩下的维度应重新归一化，仍然给出分数"


def test_rank_command_writes_four_artifacts(tmp_path):
    _fixture(tmp_path)
    out = rank_cmd.rank(out=str(tmp_path), no_collect=True, top=10)

    meta = out["meta"]
    report = Path(meta["report"])
    assert report.is_file()
    payload = json.loads(report.read_text(encoding="utf-8"))
    assert payload["weights_used"]["dimensions"]["imbalance"]["weight"] == 0.3
    assert payload["weights_source"], "必须记录权重来源，否则无法事后归因"
    assert payload["counts"]["excluded"] == 1

    assert len(meta["csv"]) == 3
    # 用 utf-8-sig 读：报告 CSV 带 BOM（为了让 Excel 不乱码）
    header = Path(meta["csv"][0]).read_text(encoding="utf-8-sig").splitlines()[0]
    assert header.split(",")[:3] == list(rank_cmd.CATEGORY_COLUMNS)[:3]

    seller_header = Path(meta["csv"][1]).read_text(encoding="utf-8-sig").splitlines()[0]
    assert "items_sampled" in seller_header, "「采到几件货」是卖家榜的核心列"

    with Path(meta["csv"][2]).open(encoding="utf-8-sig") as fh:
        rows = list(csv.DictReader(fh))
    assert {"excluded", "exclude_reason"} <= set(rows[0])


def test_weights_override_by_json_string(tmp_path):
    cfg, source = rank_cmd._store.load_weights(tmp_path, json.dumps({"dimensions": {}}))
    assert cfg == {"dimensions": {}}
    assert "命令行" in source


def test_weights_template_copied_on_first_run(tmp_path):
    """首次运行应把包内模板复制到用户目录 —— 保证「有一份可改的文件」。"""
    path = tmp_path / "weights.json"
    assert not path.exists()
    cfg, source = rank_cmd._store.load_weights(tmp_path)
    assert path.is_file()
    assert cfg["dimensions"]["imbalance"]["weight"] == 0.3
    assert "模板" in source


def test_xlsx_writes_multi_sheet_workbook(tmp_path):
    """--xlsx 应产出多 sheet 工作簿，并带概览页说明数据来源与当时权重。"""
    import pytest

    pytest.importorskip("openpyxl")
    from openpyxl import load_workbook

    _fixture(tmp_path)
    out = rank_cmd.rank(out=str(tmp_path), no_collect=True, top=5, xlsx=True)
    path = Path(out["meta"]["xlsx"])
    assert path.is_file() and path.suffix == ".xlsx"

    wb = load_workbook(path)
    assert wb.sheetnames[:3] == ["概览", "类目榜", "卖家榜"], wb.sheetnames
    assert "商品明细" in wb.sheetnames and "被过滤" in wb.sheetnames

    overview = wb["概览"]
    cells = {r[0].value: r[1].value for r in overview.iter_rows(min_row=2)}
    assert "权重·imbalance" in cells, "概览必须写明当时用的权重，否则日后无法归因"
    assert cells.get("条数·excluded") == 1

    cats = wb["类目榜"]
    assert cats.freeze_panes == "A2", "冻结首行，滚动时表头常驻"
    assert [c.value for c in cats[1]][:2] == ["rank", "score"]


def test_xlsx_without_openpyxl_says_how_to_install(tmp_path, monkeypatch):
    """缺可选依赖时必须给出可执行的安装提示，而不是裸 ImportError。"""
    import pytest

    from goofish_cli.core.errors import GoofishError

    _fixture(tmp_path)
    monkeypatch.setitem(__import__("sys").modules, "openpyxl", None)
    with pytest.raises(GoofishError) as excinfo:
        rank_cmd.rank(out=str(tmp_path), no_collect=True, xlsx=True)
    msg = str(excinfo.value)
    assert "openpyxl" in msg and "goofish-cli[excel]" in msg
